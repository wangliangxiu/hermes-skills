# Template: 大模型学习计划生成脚本
# 使用方法：运行后会在桌面生成 "大模型学习计划.xlsx"
# 如需调整学习内容，修改下方 plan 列表中的 tuple 即可

import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "大模型学习计划"

# 样式定义
header_font = Font(name="微软雅黑", bold=True, size=12, color="FFFFFF")
header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
thin_border = Border(
    left=Side(style='thin'), right=Side(style='thin'),
    top=Side(style='thin'), bottom=Side(style='thin')
)

phase_fills = {
    "第一阶段 Python基础": PatternFill(start_color="DAEEF3", end_color="DAEEF3", fill_type="solid"),
    "第二阶段 数学基础": PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid"),
    "第三阶段 机器学习入门": PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid"),
    "第四阶段 深度学习": PatternFill(start_color="EDEDED", end_color="EDEDED", fill_type="solid"),
    "第五阶段 大模型核心": PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid"),
    "第六阶段 工程与面试": PatternFill(start_color="E4DFEC", end_color="E4DFEC", fill_type="solid"),
}

headers = ["阶段", "序号", "学习内容", "完成✓", "完成日期", "备注"]
col_widths = [18, 6, 45, 10, 14, 20]

for col, w in enumerate(col_widths, 1):
    ws.column_dimensions[chr(64+col)].width = w

for col, h in enumerate(headers, 1):
    cell = ws.cell(row=1, column=col, value=h)
    cell.font, cell.fill, cell.alignment, cell.border = header_font, header_fill, header_align, thin_border

ws.row_dimensions[1].height = 30

# ===== 在此处修改学习内容 =====
plan = [
    # (阶段, 序号, 内容, 备注)
    ("第一阶段 Python基础", 1, "print输出、变量、字符串拼接", ""),
    ("第一阶段 Python基础", 2, "input交互、类型转换(int/str)", ""),
    ("第一阶段 Python基础", 3, "条件判断 if/elif/else", ""),
    ("第一阶段 Python基础", 4, "列表 list 和 for 循环", ""),
    ("第一阶段 Python基础", 5, "字典 dict 和 while 循环", ""),
    ("第一阶段 Python基础", 6, "函数 def、参数、返回值", ""),
    ("第一阶段 Python基础", 7, "文件读写 open/with", ""),
    ("第一阶段 Python基础", 8, "面向对象基础 class", ""),
    ("第一阶段 Python基础", 9, "NumPy入门：数组、矩阵运算", "数据科学基础"),
    ("第一阶段 Python基础", 10, "Pandas入门：DataFrame处理", "数据科学基础"),
    ("第一阶段 Python基础", 11, "小项目：用Python做个计算器", ""),
    ("第二阶段 数学基础", 1, "线性代数：矩阵乘法、转置", ""),
    ("第二阶段 数学基础", 2, "线性代数：特征值、特征向量", "Transformer相关"),
    ("第二阶段 数学基础", 3, "概率论：概率分布、条件概率", ""),
    ("第二阶段 数学基础", 4, "概率论：贝叶斯定理", ""),
    ("第二阶段 数学基础", 5, "微积分：导数、链式法则", ""),
    ("第二阶段 数学基础", 6, "微积分：梯度、偏导数", ""),
    ("第三阶段 机器学习入门", 1, "什么是机器学习？", ""),
    ("第三阶段 机器学习入门", 2, "线性回归（sklearn实现）", ""),
    ("第三阶段 机器学习入门", 3, "逻辑回归（分类）", ""),
    ("第三阶段 机器学习入门", 4, "损失函数、梯度下降", ""),
    ("第三阶段 机器学习入门", 5, "过拟合、正则化", ""),
    ("第三阶段 机器学习入门", 6, "决策树、随机森林", ""),
    ("第三阶段 机器学习入门", 7, "K-Means聚类", ""),
    ("第三阶段 机器学习入门", 8, "小项目：房价预测/手写数字识别", ""),
    ("第四阶段 深度学习", 1, "神经网络原理", ""),
    ("第四阶段 深度学习", 2, "PyTorch基础", "重点框架"),
    ("第四阶段 深度学习", 3, "搭建第一个神经网络", ""),
    ("第四阶段 深度学习", 4, "CNN卷积神经网络", ""),
    ("第四阶段 深度学习", 5, "RNN/LSTM", ""),
    ("第四阶段 深度学习", 6, "激活函数/BatchNorm/Dropout", ""),
    ("第四阶段 深度学习", 7, "小项目：图片分类", ""),
    ("第五阶段 大模型核心", 1, "Transformer论文精读", "大模型的灵魂"),
    ("第五阶段 大模型核心", 2, "Self-Attention", ""),
    ("第五阶段 大模型核心", 3, "Positional Encoding", ""),
    ("第五阶段 大模型核心", 4, "BERT", ""),
    ("第五阶段 大模型核心", 5, "GPT系列", ""),
    ("第五阶段 大模型核心", 6, "预训练 vs 微调", ""),
    ("第五阶段 大模型核心", 7, "RLHF", ""),
    ("第五阶段 大模型核心", 8, "LoRA/QLoRA", ""),
    ("第五阶段 大模型核心", 9, "RAG", ""),
    ("第五阶段 大模型核心", 10, "HuggingFace跑模型", ""),
    ("第五阶段 大模型核心", 11, "项目：RAG问答系统", "简历项目1"),
    ("第五阶段 大模型核心", 12, "项目：LoRA微调", "简历项目2"),
    ("第六阶段 工程与面试", 1, "模型部署vLLM/TGI", ""),
    ("第六阶段 工程与面试", 2, "推理优化GGUF/KV Cache", ""),
    ("第六阶段 工程与面试", 3, "分布式训练概念", ""),
    ("第六阶段 工程与面试", 4, "LeetCode刷题", ""),
    ("第六阶段 工程与面试", 5, "Transformer八股", ""),
    ("第六阶段 工程与面试", 6, "整理项目经验", ""),
    ("第六阶段 工程与面试", 7, "模拟面试", ""),
]

row = 2
for phase, seq, content, note in plan:
    for col, val in enumerate([phase, seq, content, "", "", note], 1):
        cell = ws.cell(row=row, column=col, value=val)
        cell.fill = phase_fills.get(phase, PatternFill())
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        cell.border = thin_border
        cell.font = Font(name="微软雅黑", size=10)
    ws.row_dimensions[row].height = 24
    row += 1

ws.freeze_panes = "A2"
ws.auto_filter.ref = f"A1:F{row-1}"

filepath = r"C:\Users\使用者\Desktop\大模型学习计划.xlsx"
wb.save(filepath)
print(f"表格已生成：{filepath}，共 {len(plan)} 项学习内容")
