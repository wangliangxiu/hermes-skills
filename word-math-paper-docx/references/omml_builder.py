# -*- coding: utf-8 -*-
"""LaTeX -> Word 原生公式(OMML) 的可复用核心。
依赖：python-docx + lxml + latex2mathml，以及 Office 的 MML2OMML.XSL。
环境：D:\pyenv\kaoyan\Scripts\python.exe（已装齐）。
"""
import copy, glob, os, re
import latex2mathml.converter as latex2mathml
from lxml import etree
from docx.shared import Pt
from docx.oxml.ns import qn


def find_xsl():
    """自动定位 Office 自带的 MathML->OMML 样式表（不同版本/安装盘符都能找到）"""
    env = os.environ.get('MML2OMML')
    if env and os.path.exists(env):
        return env
    roots = [os.environ.get('ProgramFiles', r'C:\Program Files'),
             os.environ.get('ProgramFiles(x86)', r'C:\Program Files (x86)')]
    for base in roots:
        for sub in (r'Microsoft Office\root\Office16', r'Microsoft Office\Office16',
                    r'Microsoft Office\root\Office15'):
            p = os.path.join(base, sub, 'MML2OMML.XSL')
            if os.path.exists(p):
                return p
    for base in roots:
        hits = glob.glob(os.path.join(base, 'Microsoft Office', '**', 'MML2OMML.XSL'), recursive=True)
        if hits:
            return hits[0]
    raise FileNotFoundError(
        '找不到 MML2OMML.XSL（随 Microsoft Office 安装）。先跑 bootstrap_env.py 自检；'
        '或用环境变量 MML2OMML 指向该文件。')


XSL = find_xsl()
TRANS = etree.XSLT(etree.parse(XSL))

M_NS = 'http://schemas.openxmlformats.org/officeDocument/2006/math'
M = '{%s}' % M_NS

LATEX_FIX = [
    ('\\limits', ''),
    ('\\dfrac', '\\frac'),
    ('\\displaystyle', ''),
    ('\\,', '\\;'),
    ('\\mathrm{d}', 'd'),
    ('\\text{', '\\mathrm{'),
]
BLOCK_RE = re.compile(chr(92) * 2 + r'begin' + chr(92) + r'{(pmatrix|vmatrix|bmatrix|matrix|cases)' + chr(92) + r'}(.*?)' + chr(92) * 2 + r'end' + chr(92) + r'{\1' + chr(92) + r'}', re.S)

def fix_latex(tex):
    for a, b in LATEX_FIX:
        tex = tex.replace(a, b)
    return tex

def _el(tag, **attrs):
    e = etree.Element(M + tag)
    for k, v in attrs.items():
        e.set(M + k, v)
    return e

def _omml_children(tex):
    om = TRANS(etree.fromstring(latex2mathml.convert(fix_latex(tex)))).getroot()
    return list(om)

def _matrix_omml(env, body):
    mat = _el('m')
    for row in body.split(chr(92) * 2):      # 行分隔符 -> 用 chr(92) 拼，避免转义坑
        r = _el('mr')
        for cell in row.split('&'):
            e = _el('e')
            cell = cell.strip()
            if cell:
                for ch in _omml_children(cell):
                    e.append(copy.deepcopy(ch))
            r.append(e)
        mat.append(r)
    if env in ('pmatrix', 'vmatrix', 'bmatrix', 'cases'):
        beg, end = {'pmatrix': ('(', ')'), 'vmatrix': ('|', '|'),
                    'bmatrix': ('[', ']'), 'cases': ('{', '')}[env]
        d = _el('d')
        dpr = _el('dPr')
        dpr.append(_el('begChr', val=beg))
        dpr.append(_el('endChr', val=end))
        d.append(dpr)
        e = _el('e')
        e.append(mat)
        d.append(e)
        return d
    return mat

def build_omml(tex):
    """含矩阵/分段函数的 LaTeX 片段 -> 一个 m:oMath 元素"""
    omath = _el('oMath')
    pos, parts = 0, []
    for m in BLOCK_RE.finditer(tex):
        if m.start() > pos:
            parts.append(('tex', tex[pos:m.start()], None))
        parts.append(('block', m.group(1), m.group(2)))
        pos = m.end()
    if pos < len(tex):
        parts.append(('tex', tex[pos:], None))
    for kind, a, b in parts:
        if kind == 'tex':
            if a.strip():
                for ch in _omml_children(a):
                    omath.append(copy.deepcopy(ch))
        else:
            omath.append(_matrix_omml(a, b))
    return omath

def set_font(run, ea='宋体', asc='Times New Roman', size=10.5, bold=False):
    run.font.name = asc
    run.font.size = Pt(size)
    run.font.bold = bold
    rpr = run._element.get_or_add_rPr()
    rf = rpr.find(qn('w:rFonts'))
    if rf is None:
        from docx.oxml import OxmlElement
        rf = OxmlElement('w:rFonts')
        rpr.insert(0, rf)
    rf.set(qn('w:ascii'), asc)
    rf.set(qn('w:hAnsi'), asc)
    rf.set(qn('w:eastAsia'), ea)
    return run

def add_rich(par, text, ea='宋体', size=10.5, bold=False):
    """把含 $...$ 的整段文字写进段落：$ 之间转公式，其余当普通文字"""
    for i, seg in enumerate(text.split('$')):
        if not seg:
            continue
        if i % 2 == 1:
            par._p.append(copy.deepcopy(build_omml(seg)))
        else:
            set_font(par.add_run(seg), ea=ea, size=size, bold=bold)

if __name__ == '__main__':
    from docx import Document
    d = Document()
    p = d.add_paragraph()
    add_rich(p, '设 $A=' + chr(92) + 'begin{pmatrix}2&0&1' + chr(92) * 2 + ' 3&1&x' + chr(92) * 2 + ' 4&0&5' + chr(92) + 'end{pmatrix}$。')
    d.save('demo.docx')
    print('saved demo.docx')
