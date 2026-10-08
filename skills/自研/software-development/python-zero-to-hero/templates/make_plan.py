# Template: 大模型学习计划生成脚本
# 用法：python make_plan.py
# 在桌面生成 大模型学习计划/ 文件夹，内含总表和6个分表
# 如需调整内容，修改对应阶段的 items 列表

import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
import os

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "大模型学习计划总表"

# 样式
header_font = Font(name="微软雅黑", bold=True, size=11, color="FFFFFF")
header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
thin_border = Border(
    left=Side(style='thin'), right=Side(style='thin'),
    top=Side(style='thin'), bottom=Side(style='thin')
)
body_font = Font(name="微软雅黑", size=10)
body_align = Alignment(vertical="center", wrap_text=True)
center_align = Alignment(horizontal="center", vertical="center")
done_fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
done_font = Font(name="微软雅黑", size=10, color="006100")

phase_fills = {
    "一、Python基础": PatternFill(start_color="DAEEF3", end_color="DAEEF3", fill_type="solid"),
    "二、数学基础": PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid"),
    "三、机器学习入门": PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid"),
    "四、深度学习": PatternFill(start_color="EDEDED", end_color="EDEDED", fill_type="solid"),
    "五、大模型核心": PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid"),
    "六、工程与面试": PatternFill(start_color="E4DFEC", end_color="E4DFEC", fill_type="solid"),
}

def write_header(ws, col_widths):
    headers = ["阶段", "学习内容", "完成✓", "完成日期", "备注"]
    ws.column_dimensions['A'].width = col_widths[0]
    ws.column_dimensions['B'].width = col_widths[1]
    ws.column_dimensions['C'].width = col_widths[2]
    ws.column_dimensions['D'].width = col_widths[3]
    ws.column_dimensions['E'].width = col_widths[4]
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        cell.border = thin_border
    ws.row_dimensions[1].height = 28
    ws.freeze_panes = "A2"

def write_row(ws, row, phase, content, done=False, date="", note="", fill=None):
    ws.cell(row=row, column=1, value=phase)
    ws.cell(row=row, column=2, value=content)
    c = ws.cell(row=row, column=3, value="✔️" if done else "")
    d = ws.cell(row=row, column=4, value=date)
    ws.cell(row=row, column=5, value=note)
    for col in range(1, 6):
        cell = ws.cell(row=row, column=col)
        cell.font = body_font
        cell.alignment = body_align if col == 2 else center_align
        cell.border = thin_border
        if fill:
            cell.fill = fill
    if done:
        c.font = Font(name="微软雅黑", size=14, color="006100")
        d.font = done_font
    ws.row_dimensions[row].height = 22

# 在此处定义各阶段内容...
# 完整版参考 references/learning-plan-excel.md
