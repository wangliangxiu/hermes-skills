# -*- coding: utf-8 -*-
"""PPT 程序化体检：布局dump + COM导出PNG + 文字溢出检测 + 像素白屏检测
用法: python ppt_health_check.py <pptx路径> [导出目录]
依赖: python-pptx, win32com(本机Office), PIL
输出: 控制台三部分结果 + 导出目录下的 P01.png...PNG
"""
import sys, os

PPTX = sys.argv[1] if len(sys.argv) > 1 else None
if not PPTX or not os.path.exists(PPTX):
    sys.exit('用法: python ppt_health_check.py <pptx路径> [导出目录]')
EXPORT_DIR = sys.argv[2] if len(sys.argv) > 2 else os.path.join(
    os.environ.get('TEMP', '/tmp'), 'ppt_check')


def layout_dump():
    """python-pptx 布局 dump：找尺寸不一致/错位/空run"""
    from pptx import Presentation
    prs = Presentation(PPTX)
    print(f'页面: {prs.slide_width/914400:.2f} x {prs.slide_height/914400:.2f} in, {len(prs.slides)} 页')
    for i, slide in enumerate(prs.slides, 1):
        print(f'\n===== P{i} =====')
        for sh in slide.shapes:
            t = sh.shape_type
            info = f'  [{t}] x={sh.left/914400:.2f} y={sh.top/914400:.2f} w={sh.width/914400:.2f} h={sh.height/914400:.2f}'
            if t == 13:
                info += ' [图片]'
                try:
                    info += f' 尺寸:{sh.image.size}'
                except Exception:
                    pass
            if sh.has_text_frame:
                txt = sh.text_frame.text.strip().replace('\n', '⏎')[:45]
                info += f'  文本:{txt}'
            print(info)


def com_check():
    """COM：导出PNG + 文字溢出检测（bh > usable + 4 判定）"""
    import win32com.client
    ppt = win32com.client.Dispatch('PowerPoint.Application')
    pres = ppt.Presentations.Open(PPTX, WithWindow=False)
    os.makedirs(EXPORT_DIR, exist_ok=True)
    overflow = []
    for si, slide in enumerate(pres.Slides, 1):
        for sh in slide.Shapes:
            if not sh.HasTextFrame:
                continue
            tf = sh.TextFrame
            try:
                bh = tf.TextRange.BoundHeight
            except Exception:
                continue
            usable = sh.Height - tf.MarginTop - tf.MarginBottom
            if bh > usable + 4 and tf.TextRange.Text.strip():
                overflow.append((si, tf.TextRange.Text.strip()[:30],
                                 round(bh, 1), round(usable, 1)))
        slide.Export(os.path.join(EXPORT_DIR, f'P{si:02d}.png'), 'PNG', 1280, 720)
    pres.Close()
    ppt.Quit()
    print('\n=== 文字溢出 ===')
    print('无溢出 ✓' if not overflow else overflow)


def pixel_stats():
    """PIL 像素统计：均亮/白底%/深色% -> 白屏检测与主题判断"""
    from PIL import Image
    for f in sorted(os.listdir(EXPORT_DIR)):
        if not f.endswith('.png'):
            continue
        im = Image.open(os.path.join(EXPORT_DIR, f)).convert('RGB')
        small = im.resize((64, 36))
        px = list(small.getdata())
        avg = sum(sum(p) for p in px) / (len(px) * 3)
        white = sum(1 for p in px if p[0] > 245 and p[1] > 245 and p[2] > 245) / len(px)
        dark = sum(1 for p in px if p[0] < 60 and p[1] < 60 and p[2] < 60) / len(px)
        print(f'{f}: 均亮{avg:.0f} 白底{white*100:.0f}% 深色{dark*100:.0f}%')


if __name__ == '__main__':
    layout_dump()
    com_check()
    pixel_stats()
